"""
Excel Writer Module
Appends survey respondent results as a single row to Tally.xlsx.
Preserves existing data and accurately identifies next unfilled row.
"""
import os
import openpyxl
from openpyxl.utils import column_index_from_string
import config


class ExcelWriter:
    def __init__(self, excel_path=None):
        self.excel_path = excel_path or config.EXCEL_PATH
        # Ensure target directory exists
        target_dir = os.path.dirname(self.excel_path)
        if target_dir:
            os.makedirs(target_dir, exist_ok=True)

    def get_next_respondent_info(self):
        """
        Determines the next row number and respondent number by checking
        the first row starting from DATA_START_ROW that has no answer data (Col B).
        Returns: (next_row, next_respondent_no)
        """
        if not os.path.exists(self.excel_path):
            return config.DATA_START_ROW, 1

        wb = openpyxl.load_workbook(self.excel_path, data_only=True)
        ws = wb.active

        target_row = None
        target_respondent_no = None

        # Inspect rows starting from DATA_START_ROW (3)
        for row in range(config.DATA_START_ROW, ws.max_row + 2):
            val_a = ws.cell(row=row, column=1).value
            # An answer exists if Column B (Strand) or any grid column (C-AF) is populated
            has_answers = any(ws.cell(row=row, column=c).value is not None for c in range(2, 33))

            if not has_answers:
                target_row = row
                try:
                    target_respondent_no = int(val_a) if val_a is not None else (row - 2)
                except (ValueError, TypeError):
                    target_respondent_no = row - 2
                break

        wb.close()
        return target_row or config.DATA_START_ROW, target_respondent_no or 1

    def append_respondent_data(self, respondent_no, survey_data):
        """
        Writes a full row for one respondent into Tally.xlsx.
        survey_data maps section id to detected value(s).
        Returns: (success: bool, message: str)
        """
        if not os.path.exists(self.excel_path):
            return False, f"Excel file not found at: {self.excel_path}"

        try:
            wb = openpyxl.load_workbook(self.excel_path)
            ws = wb.active

            target_row, _ = self.get_next_respondent_info()

            # Safety guard: never overwrite header rows or rows before row 3
            if target_row < config.DATA_START_ROW:
                target_row = config.DATA_START_ROW

            # Write Column A: Respondent #
            ws.cell(row=target_row, column=1, value=int(respondent_no))

            # Write each section's data
            for sec in config.SECTIONS:
                sec_id = sec["id"]
                cols = sec["cols"]
                val = survey_data.get(sec_id)

                if sec["type"] == "strand":
                    col_idx = column_index_from_string(cols[0])
                    ws.cell(row=target_row, column=col_idx, value=int(val) if val is not None else "")
                elif sec["type"] == "grid":
                    vals_list = val if isinstance(val, (list, tuple)) else []
                    for i, col_letter in enumerate(cols):
                        col_idx = column_index_from_string(col_letter)
                        item_val = vals_list[i] if i < len(vals_list) else ""
                        ws.cell(row=target_row, column=col_idx, value=int(item_val) if item_val != "" else "")

            wb.save(self.excel_path)
            wb.close()
            return True, f"Saved respondent #{respondent_no} at row {target_row} in {os.path.basename(self.excel_path)}"
        except PermissionError:
            return False, f"Permission denied! Please close {os.path.basename(self.excel_path)} in Excel."
        except Exception as e:
            return False, f"Error saving to Excel: {str(e)}"

    def list_sheets(self):
        """Lists all sheet/file targets available in excel directory."""
        files = []
        if os.path.exists(config.EXCEL_DIR):
            for f in os.listdir(config.EXCEL_DIR):
                if f.endswith('.xlsx') and not f.startswith('~$'):
                    files.append(f)
        if not files:
            files.append(os.path.basename(self.excel_path))
        return sorted(files)

    def switch_file(self, filename):
        """Switches active excel file target."""
        if not filename.endswith('.xlsx'):
            filename += '.xlsx'
        new_path = os.path.join(config.EXCEL_DIR, filename)
        if os.path.exists(new_path):
            self.excel_path = new_path
            return True, f"Switched to {filename}"
        return False, f"File {filename} not found"

    def create_new_template(self, filename):
        """Creates a brand new empty Excel file with exact AutoTali questionnaire headers."""
        if not filename.endswith('.xlsx'):
            filename += '.xlsx'
        target_path = os.path.join(config.EXCEL_DIR, filename)
        if os.path.exists(target_path):
            return False, f"File '{filename}' already exists!"

        try:
            # If default template exists, clone headers only
            if os.path.exists(config.EXCEL_PATH):
                src_wb = openpyxl.load_workbook(config.EXCEL_PATH)
                src_ws = src_wb.active
                new_wb = openpyxl.Workbook()
                new_ws = new_wb.active
                new_ws.title = src_ws.title

                # Copy header rows (Row 1 and Row 2)
                for r in range(1, config.DATA_START_ROW):
                    for c in range(1, 33):
                        new_ws.cell(row=r, column=c, value=src_ws.cell(row=r, column=c).value)
                src_wb.close()
            else:
                # Build fresh headers from scratch
                new_wb = openpyxl.Workbook()
                new_ws = new_wb.active
                new_ws.title = "Tally"
                new_ws.cell(row=1, column=1, value="Respondent #")
                new_ws.cell(row=2, column=1, value="No.")
                col_i = 2
                for sec in config.SECTIONS:
                    for c_letter in sec["cols"]:
                        new_ws.cell(row=1, column=col_i, value=sec["name"])
                        new_ws.cell(row=2, column=col_i, value=c_letter)
                        col_i += 1

            new_wb.save(target_path)
            new_wb.close()
            self.excel_path = target_path
            return True, f"Created and switched to {filename}"
        except Exception as e:
            return False, f"Failed to create new template: {str(e)}"
